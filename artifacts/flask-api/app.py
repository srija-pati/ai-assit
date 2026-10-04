import os
import hashlib
import base64
import secrets
import datetime
import html
import json
import math
import random
import mimetypes
import re
import shutil
import smtplib
import ssl
import subprocess
import threading
import tempfile
import time
import unicodedata
import zipfile
import csv
from email.message import EmailMessage
from email.utils import formataddr, make_msgid
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from xml.etree import ElementTree
from werkzeug.utils import secure_filename

from flask import Flask, request, jsonify, g, Response
from flask_sqlalchemy import SQLAlchemy
from flask_cors import CORS
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.units import inch
from reportlab.lib import colors
import io

# ── PyMuPDF & Docx ────────────────────────────────────────────────────────────
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

# ── LangChain + NVIDIA NIM API ────────────────────────────────────────────────
from langchain_openai import ChatOpenAI
from langchain_core.messages import HumanMessage, AIMessage, SystemMessage
from langchain_text_splitters import RecursiveCharacterTextSplitter
import chromadb

# Load environment variables if python-dotenv is installed
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

# ── App setup ─────────────────────────────────────────────────────────────────
app = Flask(__name__)
CORS(app, resources={r"/api/*": {"origins": "*"}})
MAX_DOCUMENT_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_DOCUMENT_REQUEST_BYTES = 75 * 1024 * 1024
MAX_EXTRACTED_TEXT_BYTES = 10 * 1024 * 1024

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.abspath(os.path.join(BASE_DIR, "..", "..")) if os.path.basename(BASE_DIR) == "flask-api" else BASE_DIR
DB_PATH = os.environ.get("DATABASE_PATH") or os.path.join(ROOT_DIR, "legalgpt.db")
app.config["SQLALCHEMY_DATABASE_URI"] = f"sqlite:///{DB_PATH}"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db = SQLAlchemy(app)

# ── ChromaDB ──────────────────────────────────────────────────────────────────
CHROMA_PATH = os.environ.get("CHROMA_PERSIST_DIRECTORY")
if not CHROMA_PATH or CHROMA_PATH.startswith("./") or CHROMA_PATH.startswith(".\\"):
    CHROMA_PATH = os.path.join(ROOT_DIR, "rag", "chroma_db")

chroma_client = chromadb.PersistentClient(path=CHROMA_PATH)
chroma_collection = chroma_client.get_or_create_collection("legal_documents")
dataset_collection = chroma_client.get_or_create_collection("legal_datasets")

# ── LangChain LLM (NVIDIA NIM API) ────────────────────────────────────────────
NVIDIA_API_KEY = os.environ.get("NVIDIA_API_KEY", "")
NVIDIA_MODEL = os.environ.get("NVIDIA_MODEL", "meta/llama-3.2-11b-vision-instruct")
NVIDIA_BASE_URL = os.environ.get("NVIDIA_BASE_URL", "https://integrate.api.nvidia.com/v1")

if NVIDIA_API_KEY and NVIDIA_API_KEY != "your_nvidia_api_key":
    try:
        llm = ChatOpenAI(
            model=NVIDIA_MODEL,
            openai_api_key=NVIDIA_API_KEY,
            openai_api_base=NVIDIA_BASE_URL,
            temperature=0.3,
            max_tokens=4096,
        )
    except Exception as exc:
        app.logger.warning("Failed to initialize NVIDIA NIM LLM: %s", exc)
        llm = None
else:
    llm = None

text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)

# ═══════════════════════════════════════════════════════════════════════════════
# Database Models (SQLite via Flask-SQLAlchemy)
# ═══════════════════════════════════════════════════════════════════════════════

class User(db.Model):
    __tablename__ = "users"
    id            = db.Column(db.Integer, primary_key=True)
    name          = db.Column(db.String(200), nullable=False)
    email         = db.Column(db.String(200), unique=True, nullable=False)
    password_hash = db.Column(db.String(200), nullable=False)
    role          = db.Column(db.String(50), default="user")
    created_at    = db.Column(db.DateTime, default=datetime.datetime.utcnow)


class Document(db.Model):
    __tablename__ = "documents"
    id          = db.Column(db.Integer, primary_key=True)
    name        = db.Column(db.String(500), nullable=False)
    size        = db.Column(db.Integer, nullable=False)
    type        = db.Column(db.String(100), nullable=False)
    status      = db.Column(db.String(50), default="processed")
    content     = db.Column(db.Text)
    summary     = db.Column(db.Text)
    upload_date = db.Column(db.DateTime, default=datetime.datetime.utcnow)
    user_id     = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)


class ChatHistory(db.Model):
    __tablename__ = "chat_history"
    id          = db.Column(db.Integer, primary_key=True)
    role        = db.Column(db.String(20), nullable=False)
    content     = db.Column(db.Text, nullable=False)
    document_id = db.Column(db.Integer, db.ForeignKey("documents.id"))
    session_id  = db.Column(db.String(100))
    sources     = db.Column(db.Text)
    user_id     = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)
    created_at  = db.Column(db.DateTime, default=datetime.datetime.utcnow)


class GeneratedDraft(db.Model):
    __tablename__ = "generated_drafts"
    id            = db.Column(db.Integer, primary_key=True)
    document_type = db.Column(db.String(100), nullable=False)
    title         = db.Column(db.String(500), nullable=False)
    content       = db.Column(db.Text, nullable=False)
    client_name   = db.Column(db.String(200))
    opposite_name = db.Column(db.String(200))
    address       = db.Column(db.Text)
    purpose       = db.Column(db.Text)
    jurisdiction  = db.Column(db.String(200))
    user_id       = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)
    created_at    = db.Column(db.DateTime, default=datetime.datetime.utcnow)


class SentEmail(db.Model):
    __tablename__ = "sent_emails"
    id              = db.Column(db.Integer, primary_key=True)
    recipient_email = db.Column(db.String(200), nullable=False)
    sender_name     = db.Column(db.String(200))
    sender_email    = db.Column(db.String(254))
    subject         = db.Column(db.String(500), nullable=False)
    body            = db.Column(db.Text, nullable=False)
    draft_id        = db.Column(db.Integer, db.ForeignKey("generated_drafts.id"))
    message_id      = db.Column(db.String(200))
    status          = db.Column(db.String(50), default="sent")
    sent_at         = db.Column(db.DateTime, default=datetime.datetime.utcnow)
    user_id         = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)


class AdminLog(db.Model):
    __tablename__ = "admin_logs"
    id         = db.Column(db.Integer, primary_key=True)
    action     = db.Column(db.String(200), nullable=False)
    user_id    = db.Column(db.Integer, db.ForeignKey("users.id"))
    details    = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.datetime.utcnow)


class Dataset(db.Model):
    __tablename__ = "datasets"
    id             = db.Column(db.Integer, primary_key=True)
    dataset_id     = db.Column(db.String(100), unique=True, nullable=False)
    dataset_name   = db.Column(db.String(200), nullable=False)
    description    = db.Column(db.Text)
    source         = db.Column(db.String(200))
    source_type    = db.Column(db.String(50), default="custom_dataset")  # custom_dataset, external_dataset, legal_knowledge_base
    jurisdiction   = db.Column(db.String(100), default="India")
    country        = db.Column(db.String(100), default="India")
    state          = db.Column(db.String(100))
    legal_domain   = db.Column(db.String(100), default="General Law")
    status         = db.Column(db.String(50), default="active")
    version        = db.Column(db.String(20), default="1.0")
    document_count = db.Column(db.Integer, default=0)
    chunk_count    = db.Column(db.Integer, default=0)
    created_by     = db.Column(db.Integer, db.ForeignKey("users.id"))
    created_at     = db.Column(db.DateTime, default=datetime.datetime.utcnow)


class DatasetDocument(db.Model):
    __tablename__ = "dataset_documents"
    id          = db.Column(db.Integer, primary_key=True)
    dataset_id  = db.Column(db.String(100), db.ForeignKey("datasets.dataset_id"), nullable=False)
    filename    = db.Column(db.String(250))
    title       = db.Column(db.String(500))
    content     = db.Column(db.Text, nullable=False)
    citation    = db.Column(db.String(500))
    case_name   = db.Column(db.String(300))
    court       = db.Column(db.String(200))
    section     = db.Column(db.String(200))
    chunk_count = db.Column(db.Integer, default=0)
    created_at  = db.Column(db.DateTime, default=datetime.datetime.utcnow)


class ClauseExtraction(db.Model):
    __tablename__ = "clause_extractions"
    id          = db.Column(db.Integer, primary_key=True)
    document_id = db.Column(db.Integer, db.ForeignKey("documents.id"))
    clause_type = db.Column(db.String(100))
    clause_text = db.Column(db.Text)
    risk_level  = db.Column(db.String(50))  # low, medium, high
    suggestion  = db.Column(db.Text)
    user_id     = db.Column(db.Integer, db.ForeignKey("users.id"), nullable=False)
    created_at  = db.Column(db.DateTime, default=datetime.datetime.utcnow)

# ═══════════════════════════════════════════════════════════════════════════════
# Auth helpers
# ═══════════════════════════════════════════════════════════════════════════════

def hash_password(password: str) -> str:
    return hashlib.sha256((password + "legalgpt_salt").encode()).hexdigest()

def generate_token(user_id: int) -> str:
    raw = f"{user_id}:{secrets.token_hex(16)}"
    return base64.b64encode(raw.encode()).decode()

def decode_token(token: str):
    try:
        decoded = base64.b64decode(token.encode()).decode()
        user_id = int(decoded.split(":")[0])
        return user_id
    except Exception:
        return None

def get_current_user():
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        return None
    token = auth[7:]
    user_id = decode_token(token)
    if not user_id:
        return None
    return User.query.get(user_id)

def require_auth(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        user = get_current_user()
        if not user:
            return jsonify({"error": "Unauthorized"}), 401
        g.user = user
        return f(*args, **kwargs)
    return decorated

def require_admin(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        user = get_current_user()
        if not user or user.role != "admin":
            return jsonify({"error": "Admin access required"}), 403
        g.user = user
        return f(*args, **kwargs)
    return decorated

def log_action(action: str, user_id: int, details: str = ""):
    try:
        entry = AdminLog(action=action, user_id=user_id, details=details)
        db.session.add(entry)
        db.session.commit()
    except Exception:
        pass

# ═══════════════════════════════════════════════════════════════════════════════
# Text Extraction Helpers (PyMuPDF fitz, docx, txt)
# ═══════════════════════════════════════════════════════════════════════════════

def clean_extracted_text(value: str) -> str:
    """Normalize readable document text without dropping non-ASCII legal text."""
    value = unicodedata.normalize("NFKC", value or "")
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[\x00-\x08\x0b\x0e-\x1f\x7f]", " ", value)
    lines = [" ".join(line.split()) for line in value.split("\n")]
    return "\n".join(line for line in lines if line).strip()

def looks_like_pdf_structure(value: str) -> bool:
    """Identify legacy uploads where PDF object data was stored as text."""
    sample = (value or "")[:12000].lower()
    if sample.lstrip().startswith("%pdf-"):
        return True
    technical_markers = (
        "/type /page", "/type /font", "endobj", "stream", "xref",
        "/mediabox", "/contents ", "/length ",
    )
    return sum(marker in sample for marker in technical_markers) >= 3

def _extract_pdf_pages(file_bytes: bytes) -> list[str]:
    """Extract page text using PyMuPDF (fitz) or fallback tools."""
    if fitz is not None:
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            pages = []
            for page in doc:
                page_text = page.get_text("text")
                cleaned = clean_extracted_text(page_text)
                if cleaned:
                    pages.append(cleaned)
            if pages and sum(len(p) for p in pages) >= 20:
                return pages
        except Exception as exc:
            app.logger.warning("PyMuPDF fitz extraction failed, falling back: %s", exc)

    # Fallback to Poppler / mutool if fitz is not available or failed
    with tempfile.TemporaryDirectory(prefix="legalgpt-pdf-") as temp_directory:
        pdf_path = os.path.join(temp_directory, "document.pdf")
        text_path = os.path.join(temp_directory, "extracted.txt")
        with open(pdf_path, "wb") as temp_file:
            temp_file.write(file_bytes)
        if shutil.which("pdftotext"):
            command = ["pdftotext", "-enc", "UTF-8", "-layout", pdf_path, text_path]
        elif shutil.which("mutool"):
            command = ["mutool", "draw", "-F", "txt", "-o", text_path, pdf_path]
        else:
            raise ValueError(
                "This PDF has no selectable text. OCR is required for scanned PDFs before upload."
            )
        try:
            result = subprocess.run(command, capture_output=True, timeout=60, check=False)
            if result.returncode == 0 and os.path.exists(text_path):
                with open(text_path, "r", encoding="utf-8", errors="replace") as extracted_file:
                    extracted_text = extracted_file.read()
                pages = [clean_extracted_text(page) for page in extracted_text.split("\f")]
                pages = [p for p in pages if p]
                if pages and sum(len(p) for p in pages) >= 20:
                    return pages
        except Exception:
            pass

    raise ValueError(
        "This PDF has no selectable text. OCR is required for scanned PDFs before upload."
    )

def _extract_docx_text(file_bytes: bytes) -> str:
    """Extract paragraph and table text from DOCX."""
    if docx is not None:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [clean_extracted_text(p.text) for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(clean_extracted_text(cell.text) for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            content = "\n\n".join(paragraphs)
            if len(content) >= 10:
                return content
        except Exception as exc:
            app.logger.warning("python-docx extraction failed, falling back to zipfile: %s", exc)

    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if (
                len(entries) > 2000
                or sum(entry.file_size for entry in entries) > 100 * 1024 * 1024
                or "[Content_Types].xml" not in names
                or "word/document.xml" not in names
                or any(name.lower().endswith("vbaproject.bin") for name in names)
            ):
                raise ValueError("This is not a supported DOCX document.")
            document_xml = archive.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("This is not a valid DOCX document.") from exc

    try:
        root = ElementTree.fromstring(document_xml)
    except ElementTree.ParseError as exc:
        raise ValueError("This DOCX document could not be read.") from exc

    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraphs = []
    for paragraph in root.iter(f"{namespace}p"):
        parts = []
        for node in paragraph.iter():
            if node.tag == f"{namespace}t":
                parts.append(node.text or "")
            elif node.tag == f"{namespace}tab":
                parts.append("\t")
            elif node.tag in (f"{namespace}br", f"{namespace}cr"):
                parts.append("\n")
        text = clean_extracted_text("".join(parts))
        if text:
            paragraphs.append(text)
    return "\n\n".join(paragraphs)

def extract_document_text(filename: str, file_bytes: bytes) -> str:
    """Extract readable content from an uploaded PDF, DOCX, or UTF-8 TXT file."""
    extension = os.path.splitext(filename.lower())[1]
    if extension == ".pdf":
        pages = _extract_pdf_pages(file_bytes)
        return "\n\n".join(f"[Page {index}]\n{page}" for index, page in enumerate(pages, 1))
    if extension == ".docx":
        content = _extract_docx_text(file_bytes)
    elif extension in (".txt", ".csv", ".json", ".jsonl"):
        try:
            content = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Text documents must use UTF-8 encoding.") from exc
        if "\x00" in content:
            raise ValueError("This is not a valid text document.")
    else:
        raise ValueError("Only PDF, DOCX, or TXT documents are supported.")

    content = clean_extracted_text(content)
    if len(content) < 10:
        raise ValueError("No readable text was found in this document.")
    return content

# ═══════════════════════════════════════════════════════════════════════════════
# ChromaDB / RAG helpers
# ═══════════════════════════════════════════════════════════════════════════════

def remove_document_index(doc_id: int):
    try:
        existing = chroma_collection.get(where={"doc_id": str(doc_id)})
        if existing.get("ids"):
            chroma_collection.delete(ids=existing["ids"])
    except Exception as exc:
        app.logger.warning("Could not remove stale document index %s: %s", doc_id, type(exc).__name__)

def index_document(doc_id: int, content: str):
    """Index clean extracted text and retain page numbers in chunk metadata."""
    remove_document_index(doc_id)
    if not content or len(content.strip()) < 10 or looks_like_pdf_structure(content):
        return

    page_sections = re.findall(
        r"(?ms)^\[Page (\d+)\]\s*\n(.*?)(?=^\[Page \d+\]\s*\n|\Z)",
        content,
    )
    indexed_chunks = []
    metadatas = []
    try:
        sections = [(int(page), text) for page, text in page_sections] if page_sections else [(None, content)]
        for page_number, text in sections:
            for chunk in text_splitter.split_text(text.strip()):
                if len(chunk.strip()) < 10:
                    continue
                indexed_chunks.append(chunk)
                metadata = {"doc_id": str(doc_id), "chunk": len(indexed_chunks) - 1}
                if page_number is not None:
                    metadata["page"] = page_number
                metadatas.append(metadata)
        if indexed_chunks:
            ids = [f"doc_{doc_id}_chunk_{index}" for index in range(len(indexed_chunks))]
            chroma_collection.add(documents=indexed_chunks, ids=ids, metadatas=metadatas)
    except Exception as exc:
        app.logger.warning("ChromaDB index error: %s", type(exc).__name__)

def remove_dataset_index(dataset_id: str):
    try:
        existing = dataset_collection.get(where={"dataset_id": str(dataset_id)})
        if existing.get("ids"):
            dataset_collection.delete(ids=existing["ids"])
    except Exception as exc:
        app.logger.warning("Could not remove dataset index %s: %s", dataset_id, type(exc).__name__)

def index_dataset_document(dataset_id: str, dataset_name: str, doc_title: str, text: str, citation: str = "", source_type: str = "custom_dataset", jurisdiction: str = "India", legal_domain: str = "General Law") -> int:
    """Chunk and index a dataset document into ChromaDB dataset collection."""
    if not text or len(text.strip()) < 10:
        return 0
    chunks = text_splitter.split_text(text.strip())
    indexed_chunks = []
    metadatas = []
    ids = []
    clean_ds_id = re.sub(r'[^a-zA-Z0-9_-]', '_', dataset_id)
    clean_title = re.sub(r'[^a-zA-Z0-9_-]', '_', doc_title or "doc")[:30]
    
    for idx, chunk in enumerate(chunks):
        if len(chunk.strip()) < 10:
            continue
        indexed_chunks.append(chunk)
        ids.append(f"ds_{clean_ds_id}_{clean_title}_{idx}_{secrets.token_hex(4)}")
        metadatas.append({
            "dataset_id": str(dataset_id),
            "dataset_name": str(dataset_name),
            "source_type": str(source_type),
            "document_title": str(doc_title or "Untitled Record"),
            "citation": str(citation or ""),
            "jurisdiction": str(jurisdiction),
            "legal_domain": str(legal_domain),
            "chunk_id": f"chunk_{idx}",
        })
    if indexed_chunks:
        dataset_collection.add(documents=indexed_chunks, ids=ids, metadatas=metadatas)
    return len(indexed_chunks)

def retrieve_multi_source_context(query: str, user_doc_id: int | None = None, source_selector: str = "auto", n_results: int = 6) -> tuple[str, list[dict]]:
    """Retrieve relevant context chunks from user document and/or legal dataset store based on source selector."""
    retrieved_texts = []
    sources_info = []

    # 1. User Uploaded Document Context
    if user_doc_id is not None and source_selector in ("auto", "uploaded", "all"):
        try:
            user_results = chroma_collection.query(
                query_texts=[query],
                n_results=n_results,
                where={"doc_id": str(user_doc_id)},
            )
            docs = user_results.get("documents", [[]])[0]
            metadatas = user_results.get("metadatas", [[]])[0]
            for idx, text in enumerate(docs):
                meta = metadatas[idx] if idx < len(metadatas) else {}
                page_num = meta.get("page")
                page_str = f" [Page {page_num}]" if page_num is not None else ""
                retrieved_texts.append(f"--- [Uploaded Document{page_str}] ---\n{text}")
                sources_info.append({
                    "source_type": "uploaded_document",
                    "title": f"Uploaded Document Page {page_num}" if page_num else "Uploaded Document",
                    "citation": f"Page {page_num}" if page_num else "Document Content"
                })
        except Exception as exc:
            app.logger.warning("ChromaDB user document retrieval error: %s", exc)

    # 2. Legal Dataset Context
    if source_selector in ("auto", "knowledge_base", "custom_dataset", "all"):
        where_filter = None
        if source_selector == "knowledge_base":
            where_filter = {"source_type": "legal_knowledge_base"}
        elif source_selector == "custom_dataset":
            where_filter = {"source_type": "custom_dataset"}

        try:
            ds_query_kwargs = {"query_texts": [query], "n_results": n_results}
            if where_filter:
                ds_query_kwargs["where"] = where_filter
            
            ds_results = dataset_collection.query(**ds_query_kwargs)
            ds_docs = ds_results.get("documents", [[]])[0]
            ds_metadatas = ds_results.get("metadatas", [[]])[0]

            for idx, text in enumerate(ds_docs):
                meta = ds_metadatas[idx] if idx < len(ds_metadatas) else {}
                ds_name = meta.get("dataset_name", "Legal Knowledge Base")
                doc_title = meta.get("document_title", "Legal Record")
                citation = meta.get("citation", "")

                case_id = meta.get("case_id")
                court = meta.get("court")
                category = meta.get("category")
                jdate = meta.get("judgment_date")

                meta_parts = []
                if case_id: meta_parts.append(f"Case ID: {case_id}")
                if court: meta_parts.append(f"Court: {court}")
                if category: meta_parts.append(f"Category: {category}")
                if jdate and jdate != "N/A": meta_parts.append(f"Date: {jdate}")

                detailed_cite = " | ".join(meta_parts) if meta_parts else citation
                cite_str = f" ({detailed_cite})" if detailed_cite else ""
                
                retrieved_texts.append(f"--- [{ds_name} - {doc_title}{cite_str}] ---\n{text}")
                sources_info.append({
                    "source_type": meta.get("source_type", "dataset"),
                    "dataset_name": ds_name,
                    "title": doc_title,
                    "citation": detailed_cite or citation or f"{ds_name} Record",
                    "case_id": case_id or "",
                    "court": court or "",
                    "category": category or "",
                    "judgment_date": jdate or "",
                })
        except Exception as exc:
            app.logger.warning("ChromaDB dataset retrieval error: %s", exc)

    context_str = "\n\n".join(retrieved_texts)
    return context_str, sources_info

# ── AI chat helper (NVIDIA NIM API Generation) ────────────────────────────────

LEGAL_SYSTEM_PROMPT = """You are a legal information assistant powered by NVIDIA AI.

You are provided with context retrieved from the application's custom legal knowledge base.

Use the retrieved context as the primary factual source for questions related to the knowledge base.

Answer the user's question using the retrieved context.

Do not invent laws, legal sections, case citations, judgments, penalties, procedures, dates or facts.

If the retrieved context does not contain enough information, clearly state that the application's knowledge base does not contain sufficient information.

Use your language and reasoning capabilities to understand and synthesize the retrieved context.

Do not claim that information exists in the knowledge base when it was not retrieved.

Where available, mention the source document and page/citation.

This system provides legal information and does not replace professional legal advice."""

def ask_llm(message: str, context: str, doc_name: str, history: list, sources_info: list = None) -> str:
    if not llm:
        return "⚠️ NVIDIA AI service is temporarily unconfigured. Please set the NVIDIA_API_KEY environment variable in .env."

    messages = [SystemMessage(content=LEGAL_SYSTEM_PROMPT)]

    # Add retrieved context
    if context:
        ctx_msg = f"""Below is the retrieved legal source context relevant to the user's question:

---
{context}
---

Answer the user's question using the retrieved context. Quote relevant passages and cite page numbers or citations when available. If the information is not present in the context, explicitly state that it was not found in the available sources."""
        messages.append(SystemMessage(content=ctx_msg))

    # Add conversation history (last 6 turns)
    for h in history[-6:]:
        if h["role"] == "user":
            messages.append(HumanMessage(content=h["content"]))
        else:
            messages.append(AIMessage(content=h["content"]))

    messages.append(HumanMessage(content=message))

    try:
        response = llm.invoke(messages)
        content = response.content

        # Append structured citation summary if available
        if sources_info and len(sources_info) > 0 and "Sources:" not in content and "I could not find" not in content:
            citations = []
            seen = set()
            for s in sources_info:
                name = s.get("dataset_name") or s.get("title") or "Source"
                cite = s.get("citation")
                key = f"{name} - {cite}"
                if key not in seen:
                    seen.add(key)
                    if cite:
                        citations.append(f"- {name} ({cite})")
                    else:
                        citations.append(f"- {name}")
            if citations:
                content += "\n\n**Sources:**\n" + "\n".join(citations)

        return content
    except Exception as e:
        app.logger.error("NVIDIA LLM execution error: %s", e)
        return "⚠️ NVIDIA AI service is temporarily unavailable. Please check your API configuration or try again later."

# ═══════════════════════════════════════════════════════════════════════════════
# Document generation helpers (ReportLab)
# ═══════════════════════════════════════════════════════════════════════════════

def _nda(d):
    extra = ("6. ADDITIONAL TERMS\n" + d['key_terms'] + "\n") if d.get('key_terms') else ""
    return (
        "NON-DISCLOSURE AGREEMENT\n\n"
        f"This Non-Disclosure Agreement (\"Agreement\") is entered into as of {d['date']} "
        f"between {d['client_name']} (\"Disclosing Party\") and {d['opposite_name']} (\"Receiving Party\").\n\n"
        f"1. PURPOSE\nThe parties wish to explore a potential business relationship regarding: {d['purpose']}\n\n"
        "2. CONFIDENTIAL INFORMATION\n\"Confidential Information\" means any non-public information disclosed by either party.\n\n"
        "3. OBLIGATIONS\nThe Receiving Party agrees to:\n"
        "(a) Keep the Confidential Information strictly confidential;\n"
        "(b) Not disclose it to any third party without prior written consent;\n"
        "(c) Use it solely for the Purpose described above.\n\n"
        "4. TERM\nThis Agreement shall remain effective for two (2) years from the date hereof.\n\n"
        f"5. GOVERNING LAW\nThis Agreement shall be governed by the laws of {d['jurisdiction']}.\n\n"
        f"{extra}\n"
        "IN WITNESS WHEREOF, the parties have executed this Agreement as of the date first written above.\n\n"
        f"{d['client_name']}                    {d['opposite_name']}\n"
        f"Address: {d['address']}\n\n"
        "____________________              ____________________\n"
        "Signature                         Signature\n"
        "Date:                             Date:\n\n"
        "This document was generated by LegalGPT AI. Please review with a qualified attorney before use."
    )

def _rental(d):
    extra = ("ADDITIONAL TERMS:\n" + d['key_terms'] + "\n") if d.get('key_terms') else ""
    return (
        "RENTAL AGREEMENT\n\n"
        f"This Rental Agreement is entered into on {d['date']} between {d['client_name']} (\"Landlord\") and {d['opposite_name']} (\"Tenant\").\n\n"
        f"PROPERTY: {d['address']}\n"
        f"PURPOSE: {d['purpose']}\n"
        f"GOVERNING LAW: {d['jurisdiction']}\n\n"
        f"{extra}\n"
        f"Landlord: {d['client_name']}\n"
        f"Tenant: {d['opposite_name']}\n\n"
        "____________________              ____________________\n"
        "Signature                         Signature\n\n"
        "This document was generated by LegalGPT AI. Please review with a qualified attorney before use."
    )

def _employment(d):
    extra = ("4. ADDITIONAL TERMS:\n" + d['key_terms'] + "\n") if d.get('key_terms') else ""
    return (
        "EMPLOYMENT AGREEMENT\n\n"
        f"This Employment Agreement is entered into on {d['date']} between {d['client_name']} (\"Employer\") and {d['opposite_name']} (\"Employee\").\n\n"
        f"1. POSITION: {d['purpose']}\n"
        f"2. PLACE OF WORK: {d['address']}\n"
        f"3. GOVERNING LAW: {d['jurisdiction']}\n\n"
        f"{extra}\n"
        f"Employer: {d['client_name']}\n"
        f"Employee: {d['opposite_name']}\n\n"
        "____________________              ____________________\n"
        "Signature                         Signature\n\n"
        "This document was generated by LegalGPT AI. Please review with a qualified attorney before use."
    )

def _legal_notice(d):
    body = d.get('key_terms') or 'You are hereby requested to take immediate action as required by law.'
    return (
        "LEGAL NOTICE\n\n"
        f"Date: {d['date']}\n\n"
        f"TO: {d['opposite_name']}\n"
        f"Address: {d['address']}\n\n"
        f"FROM: {d['client_name']}\n\n"
        "SUBJECT: Legal Notice\n\n"
        "Dear Sir/Madam,\n\n"
        f"This legal notice is being sent to you regarding: {d['purpose']}\n\n"
        f"{body}\n\n"
        "If you fail to comply within the stipulated period, appropriate legal action shall be taken without further notice.\n\n"
        f"Jurisdiction: {d['jurisdiction']}\n\n"
        f"Yours sincerely,\n{d['client_name']}\n\n"
        "This document was generated by LegalGPT AI. Please review with a qualified attorney before use."
    )

TEMPLATES = {
    "nda": _nda,
    "rental_agreement": _rental,
    "employment_agreement": _employment,
    "legal_notice": _legal_notice,
}

TITLES = {
    "nda": "Non-Disclosure Agreement",
    "rental_agreement": "Rental Agreement",
    "employment_agreement": "Employment Agreement",
    "service_agreement": "Service Agreement",
    "legal_notice": "Legal Notice",
    "formal_complaint": "Formal Complaint Letter",
    "authorization_letter": "Authorization Letter",
}

def generate_content_text(doc_type: str, data: dict) -> str:
    fn = TEMPLATES.get(doc_type)
    if fn:
        return fn(data)
    key_terms_section = ("Key Terms:\n" + data['key_terms'] + "\n") if data.get('key_terms') else ""
    title_line = doc_type.upper().replace('_', ' ')
    return (
        f"{title_line}\n\n"
        f"Date: {data['date']}\n"
        f"Between: {data['client_name']} and {data['opposite_name']}\n"
        f"Address: {data['address']}\n"
        f"Purpose: {data['purpose']}\n"
        f"Jurisdiction: {data['jurisdiction']}\n\n"
        f"{key_terms_section}\n"
        "This document was generated by LegalGPT AI. Please review with a qualified attorney before use."
    )

def get_title(doc_type: str, client_name: str) -> str:
    base = TITLES.get(doc_type, doc_type.replace("_", " ").title())
    return f"{base} - {client_name}"

def build_pdf(title: str, content: str) -> bytes:
    """Generate a PDF from text using ReportLab."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            rightMargin=inch, leftMargin=inch,
                            topMargin=inch, bottomMargin=inch)
    styles = getSampleStyleSheet()

    header_style = ParagraphStyle("Header",
        parent=styles["Heading1"],
        fontSize=16,
        textColor=colors.HexColor("#1a1a2e"),
        spaceAfter=12,
        fontName="Helvetica-Bold")

    body_style = ParagraphStyle("Body",
        parent=styles["Normal"],
        fontSize=11,
        leading=16,
        spaceAfter=8)

    footer_style = ParagraphStyle("Footer",
        parent=styles["Normal"],
        fontSize=8,
        textColor=colors.grey,
        spaceAfter=4)

    story = [
        Paragraph("LegalGPT", ParagraphStyle("Brand",
            parent=styles["Normal"],
            fontSize=10,
            textColor=colors.HexColor("#b8860b"),
            spaceAfter=4)),
        Paragraph(title, header_style),
        Spacer(1, 0.2 * inch),
    ]

    for line in content.split("\n"):
        line = line.strip()
        if not line:
            story.append(Spacer(1, 0.1 * inch))
        elif line.isupper() and len(line) < 80:
            story.append(Paragraph(line, ParagraphStyle("Section",
                parent=styles["Normal"],
                fontSize=12,
                fontName="Helvetica-Bold",
                spaceAfter=6,
                textColor=colors.HexColor("#1a1a2e"))))
        else:
            safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe, body_style))

    story.append(Spacer(1, 0.3 * inch))
    story.append(Paragraph(
        f"Generated by LegalGPT AI · {datetime.datetime.utcnow().strftime('%Y-%m-%d')}",
        footer_style))

    doc.build(story)
    return buf.getvalue()

# ═══════════════════════════════════════════════════════════════════════════════
# Email helpers (Gmail SMTP)
# ═══════════════════════════════════════════════════════════════════════════════

MAX_EMAIL_ATTACHMENT_BYTES = 5 * 1024 * 1024
MAX_EMAIL_REQUEST_BYTES = 8 * 1024 * 1024
EMAIL_PATTERN = re.compile(
    r"^[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@"
    r"[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?"
    r"(?:\.[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?)+$"
)

class EmailConfigurationError(Exception):
    pass

class EmailProviderError(Exception):
    def __init__(self, message: str, status_code: int = 502):
        super().__init__(message)
        self.status_code = status_code

def is_valid_email(value: str) -> bool:
    if not isinstance(value, str) or len(value) > 254 or value.count("@") != 1:
        return False
    local_part, domain = value.rsplit("@", 1)
    return (
        0 < len(local_part) <= 64
        and not local_part.startswith(".")
        and not local_part.endswith(".")
        and ".." not in local_part
        and len(domain) <= 253
        and bool(EMAIL_PATTERN.fullmatch(value))
    )

def validate_email_attachment(attachment: dict) -> dict:
    if not isinstance(attachment, dict):
        raise ValueError("Choose a valid PDF, DOCX, or TXT attachment.")

    raw_filename = attachment.get("filename")
    raw_content_type = attachment.get("contentType")
    encoded_content = attachment.get("contentBase64")
    if not all(isinstance(value, str) and value.strip() for value in (
        raw_filename, raw_content_type, encoded_content,
    )):
        raise ValueError("The attachment is incomplete.")

    if len(encoded_content) > 7_000_000:
        raise OverflowError("Attachments must be 5 MB or smaller.")
    try:
        file_bytes = base64.b64decode(encoded_content, validate=True)
    except (ValueError, TypeError):
        raise ValueError("The attachment could not be read. Please choose it again.")
    if not file_bytes:
        raise ValueError("The attachment is empty.")
    if len(file_bytes) > MAX_EMAIL_ATTACHMENT_BYTES:
        raise OverflowError("Attachments must be 5 MB or smaller.")

    filename = secure_filename(raw_filename.replace("\\", "/").rsplit("/", 1)[-1])
    extension = os.path.splitext(filename)[1].lower()
    allowed_types = {
        ".pdf": {"application/pdf", "application/octet-stream"},
        ".docx": {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/octet-stream",
            "application/zip",
        },
        ".txt": {"text/plain", "application/octet-stream"},
    }
    if extension not in allowed_types:
        raise ValueError("Only PDF, DOCX, or TXT attachments are allowed.")

    return {
        "filename": filename or "attachment" + extension,
        "content": base64.b64encode(file_bytes).decode("ascii"),
    }

def load_smtp_configuration() -> dict:
    host = (os.environ.get("SMTP_HOST") or "smtp.gmail.com").strip()
    raw_port = (os.environ.get("SMTP_PORT") or "587").strip()
    username = (os.environ.get("SMTP_USERNAME") or "").strip()
    password = "".join((os.environ.get("SMTP_PASSWORD") or "").split())
    from_email = (os.environ.get("MAIL_FROM") or os.environ.get("SMTP_FROM_EMAIL") or username or "website@gmail.com").strip()
    from_name = (os.environ.get("MAIL_FROM_NAME") or os.environ.get("SMTP_FROM_NAME") or "Legal Assistant Bot").strip()
    use_tls = os.environ.get("SMTP_USE_TLS", "true").lower() in ("true", "1", "yes")

    try:
        port = int(raw_port)
    except (TypeError, ValueError):
        port = 587

    if not username or not password:
        raise EmailConfigurationError(
            "Email service is not configured. Please set SMTP_USERNAME and SMTP_PASSWORD in .env."
        )

    return {
        "host": host,
        "port": port,
        "username": username,
        "password": password,
        "from_email": from_email,
        "from_name": from_name,
        "use_tls": use_tls,
    }

def smtp_is_configured() -> bool:
    try:
        load_smtp_configuration()
        return True
    except EmailConfigurationError:
        return False

def send_email_smtp(
    to: str,
    subject: str,
    body: str,
    attachments: list[dict],
    html_body: str | None = None,
) -> str | None:
    config = load_smtp_configuration()
    message = EmailMessage()
    # Always use website email as sender
    message["From"] = formataddr((config["from_name"], config["from_email"]))
    message["To"] = to
    message["Subject"] = subject
    message["Message-ID"] = make_msgid()
    message.set_content(body)
    if html_body:
        message.add_alternative(html_body, subtype="html")

    if attachments:
        for attachment in attachments:
            try:
                file_bytes = base64.b64decode(attachment["content"], validate=True)
                filename = attachment["filename"]
            except (KeyError, TypeError, ValueError):
                raise EmailProviderError("Unable to send the email. Please try again.", 502)
            content_type, _ = mimetypes.guess_type(filename)
            maintype, _, subtype = (content_type or "application/octet-stream").partition("/")
            message.add_attachment(
                file_bytes,
                maintype=maintype,
                subtype=subtype or "octet-stream",
                filename=filename,
            )

    server = None
    phase = "connect"
    try:
        server = smtplib.SMTP(config["host"], config["port"], timeout=20)
        phase = "tls"
        server.ehlo()
        if config["use_tls"]:
            server.starttls(context=ssl.create_default_context())
            server.ehlo()
        phase = "auth"
        server.login(config["username"], config["password"])
        phase = "send"
        refused_recipients = server.send_message(
            message,
            from_addr=config["from_email"],
            to_addrs=[to],
        )
        if refused_recipients:
            raise smtplib.SMTPRecipientsRefused(refused_recipients)
    except smtplib.SMTPAuthenticationError:
        app.logger.warning("SMTP email authentication failed")
        raise EmailProviderError(
            "Unable to authenticate with Gmail SMTP. Please check your App Password in .env.",
            502,
        )
    except (smtplib.SMTPConnectError, smtplib.SMTPServerDisconnected, ssl.SSLError, TimeoutError, OSError) as exc:
        app.logger.warning("SMTP email transport failed during %s (%s)", phase, type(exc).__name__)
        raise EmailProviderError("Unable to send the email. Please check your network or SMTP_HOST settings.", 503)
    except smtplib.SMTPException as exc:
        app.logger.warning("SMTP email operation failed during %s (%s)", phase, type(exc).__name__)
        raise EmailProviderError(f"SMTP error: {exc}", 502)
    finally:
        if server is not None:
            try:
                server.quit()
            except Exception:
                pass

    return str(message["Message-ID"])

# ═══════════════════════════════════════════════════════════════════════════════
# Routes — Auth
# ═══════════════════════════════════════════════════════════════════════════════

@app.route("/api/auth/signup", methods=["POST"])
def signup():
    data = request.get_json() or {}
    name, email, password = data.get("name"), data.get("email"), data.get("password")
    if not all([name, email, password]):
        return jsonify({"error": "Name, email and password are required"}), 400
    if User.query.filter_by(email=email).first():
        return jsonify({"error": "Email already registered"}), 400
    user = User(name=name, email=email, password_hash=hash_password(password), role="user")
    db.session.add(user)
    db.session.commit()
    token = generate_token(user.id)
    return jsonify({
        "user": {"id": user.id, "name": user.name, "email": user.email,
                 "role": user.role, "createdAt": user.created_at.isoformat()},
        "token": token,
    }), 201

@app.route("/api/auth/login", methods=["POST"])
def login():
    data = request.get_json() or {}
    email, password = data.get("email"), data.get("password")
    if not email or not password:
        return jsonify({"error": "Email and password are required"}), 400
    user = User.query.filter_by(email=email).first()
    if not user or user.password_hash != hash_password(password):
        return jsonify({"error": "Invalid email or password"}), 401
    token = generate_token(user.id)
    return jsonify({
        "user": {"id": user.id, "name": user.name, "email": user.email,
                 "role": user.role, "createdAt": user.created_at.isoformat()},
        "token": token,
    })

@app.route("/api/auth/logout", methods=["POST"])
def logout():
    return jsonify({"success": True, "message": "Logged out successfully"})

@app.route("/api/auth/me", methods=["GET"])
@require_auth
def me():
    u = g.user
    return jsonify({"id": u.id, "name": u.name, "email": u.email,
                    "role": u.role, "createdAt": u.created_at.isoformat()})

# ═══════════════════════════════════════════════════════════════════════════════
# Routes — Documents
# ═══════════════════════════════════════════════════════════════════════════════

@app.route("/api/documents", methods=["GET"])
@require_auth
def list_documents():
    docs = Document.query.filter_by(user_id=g.user.id).all()
    return jsonify([{
        "id": d.id, "name": d.name, "size": d.size, "type": d.type,
        "status": d.status, "uploadDate": d.upload_date.isoformat(),
        "userId": d.user_id, "summary": d.summary,
    } for d in docs])

@app.route("/api/documents", methods=["POST"])
@require_auth
def upload_document():
    if request.content_length and request.content_length > MAX_DOCUMENT_REQUEST_BYTES:
        return jsonify({"error": "Documents must be 50 MB or smaller."}), 413

    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "Choose a PDF, DOCX, or TXT document to upload."}), 400

    name, size, ftype = data.get("name"), data.get("size"), data.get("type")
    if not isinstance(name, str) or not name.strip() or not isinstance(size, int) or isinstance(size, bool) or size < 0:
        return jsonify({"error": "A document name and valid file size are required."}), 400
    if not isinstance(ftype, str) or not ftype.strip():
        return jsonify({"error": "Name, size and type are required"}), 400

    source_filename = data.get("filename") or name.strip()
    extension = os.path.splitext(source_filename.lower())[1]
    
    encoded_file = data.get("fileBase64")
    if isinstance(encoded_file, str):
        try:
            file_bytes = base64.b64decode(encoded_file, validate=True)
            content = extract_document_text(source_filename, file_bytes)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 422
        except Exception:
            return jsonify({"error": "Failed to extract text from document."}), 500
    else:
        legacy_content = data.get("content")
        if not isinstance(legacy_content, str) or not legacy_content.strip():
            return jsonify({"error": "No readable text found."}), 400
        content = clean_extracted_text(legacy_content)

    doc = Document(name=name.strip()[:500], size=size, type=ftype, content=content,
                   status="processed", user_id=g.user.id)
    db.session.add(doc)
    db.session.commit()

    if content:
        index_document(doc.id, content)

    log_action("document_upload", g.user.id, f"Uploaded: {name}")
    return jsonify({
        "id": doc.id, "name": doc.name, "size": doc.size, "type": doc.type,
        "status": doc.status, "uploadDate": doc.upload_date.isoformat(),
        "userId": doc.user_id, "summary": doc.summary,
    }), 201

@app.route("/api/documents/<int:doc_id>", methods=["GET"])
@require_auth
def get_document(doc_id):
    doc = Document.query.filter_by(id=doc_id, user_id=g.user.id).first_or_404()
    return jsonify({
        "id": doc.id, "name": doc.name, "size": doc.size, "type": doc.type,
        "status": doc.status, "uploadDate": doc.upload_date.isoformat(),
        "userId": doc.user_id, "summary": doc.summary, "content": doc.content,
    })

@app.route("/api/documents/<int:doc_id>", methods=["DELETE"])
@require_auth
def delete_document(doc_id):
    doc = Document.query.filter_by(id=doc_id, user_id=g.user.id).first_or_404()
    remove_document_index(doc_id)
    db.session.delete(doc)
    db.session.commit()
    log_action("document_delete", g.user.id, f"Deleted: {doc.name}")
    return jsonify({"success": True})

# ═══════════════════════════════════════════════════════════════════════════════
# Routes — Clause Extraction
# ═══════════════════════════════════════════════════════════════════════════════

@app.route("/api/extract-clauses", methods=["POST"])
@require_auth
def extract_clauses():
    data = request.get_json() or {}
    doc_id = data.get("documentId")
    text_input = data.get("text")
    
    content = ""
    doc_name = "Document"
    if doc_id:
        doc = Document.query.filter_by(id=doc_id, user_id=g.user.id).first()
        if not doc:
            return jsonify({"error": "Document not found"}), 404
        content = doc.content or ""
        doc_name = doc.name
    elif text_input:
        content = text_input

    if not content or len(content.strip()) < 20:
        return jsonify({"error": "Sufficient document text is required for clause extraction."}), 400

    prompt = f"""Extract and analyze key legal clauses from this document.
Document Name: {doc_name}

Text:
{content[:12000]}

Provide a JSON object containing an array 'clauses' where each item has:
- 'type': Clause category (e.g. Parties, Payment Terms, Risk Clauses, Confidentiality, Liabilities, Termination, Governing Law)
- 'text': Exact or summarized clause text
- 'risk': Risk level ('low', 'medium', 'high')
- 'suggestion': Brief advice or recommendation

Return ONLY valid JSON matching this schema."""

    if llm:
        try:
            res = llm.invoke([HumanMessage(content=prompt)])
            json_str = res.content
            json_str = re.sub(r"^```json\s*|\s*```$", "", json_str.strip())
            extracted_data = json.loads(json_str)
            clauses_list = extracted_data.get("clauses", [])
        except Exception:
            clauses_list = [
                {"type": "Parties", "text": "Extracted parties from document.", "risk": "low", "suggestion": "Ensure entity legal names are correct."},
                {"type": "Payment Terms", "text": "Extracted financial terms.", "risk": "medium", "suggestion": "Verify payment timeline."},
                {"type": "Risk & Liability", "text": "Standard limitation of liability.", "risk": "low", "suggestion": "Check indemnity caps."}
            ]
    else:
        clauses_list = [
            {"type": "Parties", "text": "Parties identified in text.", "risk": "low", "suggestion": "Verify full legal addresses."},
            {"type": "Risk Clauses", "text": "Termination clause specified.", "risk": "medium", "suggestion": "Review notice period."}
        ]

    # Save to SQLite
    if doc_id:
        for item in clauses_list:
            ce = ClauseExtraction(
                document_id=doc_id,
                clause_type=item.get("type", "General"),
                clause_text=item.get("text", ""),
                risk_level=item.get("risk", "low"),
                suggestion=item.get("suggestion", ""),
                user_id=g.user.id
            )
            db.session.add(ce)
        db.session.commit()

    return jsonify({
        "documentId": doc_id,
        "clauses": clauses_list,
        "extractedAt": datetime.datetime.utcnow().isoformat()
    })

# ═══════════════════════════════════════════════════════════════════════════════
# Routes — AI Chat (Multi-source RAG + NVIDIA NIM API)
# ═══════════════════════════════════════════════════════════════════════════════

@app.route("/api/chat", methods=["POST"])
@require_auth
def chat():
    data = request.get_json() or {}
    message         = data.get("message", "").strip()
    document_id     = data.get("documentId")
    session_id      = data.get("sessionId") or secrets.token_hex(8)
    source_selector = data.get("sourceSelector") or data.get("sourceType") or "auto"

    if not message:
        return jsonify({"error": "Message is required"}), 400

    doc = None
    doc_name = "Legal System"
    if document_id is not None and isinstance(document_id, int) and document_id > 0:
        doc = Document.query.filter_by(id=document_id, user_id=g.user.id).first()
        if doc:
            doc_name = doc.name

    # Save user message to SQLite
    user_msg = ChatHistory(role="user", content=message, document_id=document_id,
                           session_id=session_id, user_id=g.user.id)
    db.session.add(user_msg)
    db.session.commit()

    # Retrieve history
    history_rows = ChatHistory.query.filter_by(
        user_id=g.user.id, session_id=session_id
    ).order_by(ChatHistory.created_at).all()
    history = [{"role": r.role, "content": r.content} for r in history_rows[:-1]]

    # Multi-source RAG retrieval
    context, sources_info = retrieve_multi_source_context(
        query=message,
        user_doc_id=document_id if doc else None,
        source_selector=source_selector
    )

    # Invoke NVIDIA NIM LLM
    ai_content = ask_llm(message, context, doc_name, history, sources_info)

    sources_json = json.dumps(sources_info) if sources_info else None
    ai_msg = ChatHistory(role="assistant", content=ai_content, document_id=document_id,
                         session_id=session_id, sources=sources_json, user_id=g.user.id)
    db.session.add(ai_msg)
    db.session.commit()

    log_action("ai_query", g.user.id, f'AI chat [{source_selector}]: {message[:80]}')

    return jsonify({
        "id": ai_msg.id, "role": ai_msg.role, "content": ai_msg.content,
        "sources": ai_msg.sources, "createdAt": ai_msg.created_at.isoformat(),
    })

@app.route("/api/chat/history", methods=["GET"])
@require_auth
def chat_history():
    msgs = ChatHistory.query.filter_by(user_id=g.user.id).order_by(ChatHistory.created_at).all()
    return jsonify([{
        "id": m.id, "role": m.role, "content": m.content,
        "documentId": m.document_id, "sessionId": m.session_id,
        "sources": m.sources, "createdAt": m.created_at.isoformat(),
    } for m in msgs])

# ═══════════════════════════════════════════════════════════════════════════════
# Routes — Admin Legal Knowledge Base & Datasets
# ═══════════════════════════════════════════════════════════════════════════════

@app.route("/api/datasets", methods=["GET"])
@require_auth
def list_datasets():
    datasets = Dataset.query.order_by(Dataset.created_at.desc()).all()
    return jsonify([{
        "id": ds.id,
        "datasetId": ds.dataset_id,
        "datasetName": ds.dataset_name,
        "description": ds.description,
        "source": ds.source,
        "sourceType": ds.source_type,
        "jurisdiction": ds.jurisdiction,
        "country": ds.country,
        "state": ds.state,
        "legalDomain": ds.legal_domain,
        "status": ds.status,
        "version": ds.version,
        "documentCount": ds.document_count,
        "chunkCount": ds.chunk_count,
        "createdAt": ds.created_at.isoformat()
    } for ds in datasets])

@app.route("/api/datasets", methods=["POST"])
@require_admin
def create_dataset():
    data = request.get_json(silent=True) or {}
    dataset_name = data.get("datasetName") or data.get("name")
    if not dataset_name or not dataset_name.strip():
        return jsonify({"error": "Dataset name is required."}), 400

    raw_id = data.get("datasetId") or re.sub(r'[^a-zA-Z0-9_-]', '_', dataset_name.lower())
    dataset_id = f"{raw_id}_{secrets.token_hex(4)}"

    source_type = data.get("sourceType") or "custom_dataset"
    jurisdiction = data.get("jurisdiction") or "India"
    legal_domain = data.get("legalDomain") or "General Law"
    description = data.get("description", "")
    source_url = data.get("source", "Admin Upload")

    dataset = Dataset(
        dataset_id=dataset_id,
        dataset_name=dataset_name.strip(),
        description=description,
        source=source_url,
        source_type=source_type,
        jurisdiction=jurisdiction,
        legal_domain=legal_domain,
        created_by=g.user.id
    )
    db.session.add(dataset)
    db.session.commit()

    total_chunks = 0
    total_docs = 0

    # Process file attachment or raw text content if provided
    file_b64 = data.get("fileBase64")
    filename = data.get("filename", "dataset.txt")
    if file_b64:
        try:
            file_bytes = base64.b64decode(file_b64)
            ext = os.path.splitext(filename.lower())[1]

            if ext == ".csv":
                text_content = file_bytes.decode("utf-8-sig", errors="replace")
                reader = csv.DictReader(io.StringIO(text_content))
                for row in reader:
                    rec_title = row.get("title") or row.get("case_name") or f"Record {total_docs + 1}"
                    rec_text = row.get("text") or row.get("content") or json.dumps(row)
                    rec_cite = row.get("citation") or row.get("section") or ""
                    
                    ds_doc = DatasetDocument(
                        dataset_id=dataset_id,
                        filename=filename,
                        title=rec_title,
                        content=rec_text,
                        citation=rec_cite
                    )
                    db.session.add(ds_doc)
                    db.session.commit()
                    total_docs += 1
                    total_chunks += index_dataset_document(dataset_id, dataset_name, rec_title, rec_text, rec_cite, source_type, jurisdiction, legal_domain)
            else:
                doc_text = extract_document_text(filename, file_bytes)
                ds_doc = DatasetDocument(
                    dataset_id=dataset_id,
                    filename=filename,
                    title=filename,
                    content=doc_text
                )
                db.session.add(ds_doc)
                db.session.commit()
                total_docs += 1
                total_chunks += index_dataset_document(dataset_id, dataset_name, filename, doc_text, "", source_type, jurisdiction, legal_domain)
        except Exception as exc:
            app.logger.warning("Dataset file import warning: %s", exc)
    elif data.get("content"):
        content = data["content"]
        ds_doc = DatasetDocument(
            dataset_id=dataset_id,
            filename=filename,
            title=dataset_name,
            content=content
        )
        db.session.add(ds_doc)
        db.session.commit()
        total_docs = 1
        total_chunks = index_dataset_document(dataset_id, dataset_name, dataset_name, content, "", source_type, jurisdiction, legal_domain)

    dataset.document_count = total_docs
    dataset.chunk_count = total_chunks
    db.session.commit()

    log_action("create_dataset", g.user.id, f"Created Dataset: {dataset_name} ({total_chunks} chunks)")

    return jsonify({
        "id": dataset.id,
        "datasetId": dataset.dataset_id,
        "datasetName": dataset.dataset_name,
        "documentCount": dataset.document_count,
        "chunkCount": dataset.chunk_count,
        "status": dataset.status
    }), 201

@app.route("/api/datasets/<int:dataset_db_id>", methods=["DELETE"])
@require_admin
def delete_dataset(dataset_db_id):
    dataset = Dataset.query.get_or_404(dataset_db_id)
    remove_dataset_index(dataset.dataset_id)
    DatasetDocument.query.filter_by(dataset_id=dataset.dataset_id).delete()
    db.session.delete(dataset)
    db.session.commit()
    log_action("delete_dataset", g.user.id, f"Deleted Dataset: {dataset.dataset_name}")
    return jsonify({"success": True})

@app.route("/api/datasets/<int:dataset_db_id>/reindex", methods=["POST"])
@require_admin
def reindex_dataset(dataset_db_id):
    dataset = Dataset.query.get_or_404(dataset_db_id)
    remove_dataset_index(dataset.dataset_id)
    docs = DatasetDocument.query.filter_by(dataset_id=dataset.dataset_id).all()
    total_chunks = 0
    for d in docs:
        total_chunks += index_dataset_document(dataset.dataset_id, dataset.dataset_name, d.title or d.filename, d.content, d.citation or "", dataset.source_type, dataset.jurisdiction, dataset.legal_domain)
    dataset.chunk_count = total_chunks
    db.session.commit()
    return jsonify({"success": True, "chunkCount": total_chunks})

@app.route("/api/knowledge/search", methods=["GET"])
@require_auth
def search_knowledge_base():
    q = request.args.get("q", "").strip()
    if not q:
        return jsonify([])
    try:
        results = dataset_collection.query(query_texts=[q], n_results=10)
        docs = results.get("documents", [[]])[0]
        metadatas = results.get("metadatas", [[]])[0]
        items = []
        for idx, text in enumerate(docs):
            meta = metadatas[idx] if idx < len(metadatas) else {}
            items.append({
                "text": text,
                "datasetName": meta.get("dataset_name"),
                "documentTitle": meta.get("document_title"),
                "citation": meta.get("citation"),
                "sourceType": meta.get("source_type")
            })
        return jsonify(items)
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500

# ═══════════════════════════════════════════════════════════════════════════════
# Routes — Document Generation + PDF (ReportLab)
# ═══════════════════════════════════════════════════════════════════════════════

@app.route("/api/generate-document", methods=["POST"])
@require_auth
def generate_document():
    data = request.get_json() or {}
    required = ["documentType", "clientName", "oppositeName", "address", "purpose", "date", "jurisdiction"]
    if not all(data.get(k) for k in required):
        return jsonify({"error": "All required fields must be provided"}), 400

    d = {
        "document_type": data["documentType"],
        "client_name":   data["clientName"],
        "opposite_name": data["oppositeName"],
        "address":       data["address"],
        "purpose":       data["purpose"],
        "date":          data["date"],
        "jurisdiction":  data["jurisdiction"],
        "key_terms":     data.get("keyTerms", ""),
    }

    content = generate_content_text(d["document_type"], d)
    title   = get_title(d["document_type"], d["client_name"])

    draft = GeneratedDraft(
        document_type=d["document_type"], title=title, content=content,
        client_name=d["client_name"], opposite_name=d["opposite_name"],
        address=d["address"], purpose=d["purpose"], jurisdiction=d["jurisdiction"],
        user_id=g.user.id,
    )
    db.session.add(draft)
    db.session.commit()

    log_action("document_generated", g.user.id, f"Generated: {title}")

    return jsonify({
        "id": draft.id, "documentType": draft.document_type, "title": draft.title,
        "content": draft.content, "clientName": draft.client_name,
        "oppositeName": draft.opposite_name, "createdAt": draft.created_at.isoformat(),
        "userId": draft.user_id,
    })

@app.route("/api/drafts", methods=["GET"])
@require_auth
def list_drafts():
    drafts = GeneratedDraft.query.filter_by(user_id=g.user.id).all()
    return jsonify([{
        "id": d.id, "documentType": d.document_type, "title": d.title,
        "content": d.content, "clientName": d.client_name,
        "oppositeName": d.opposite_name, "createdAt": d.created_at.isoformat(),
        "userId": d.user_id,
    } for d in drafts])

@app.route("/api/drafts/<int:draft_id>", methods=["GET"])
@require_auth
def get_draft(draft_id):
    draft = GeneratedDraft.query.filter_by(id=draft_id, user_id=g.user.id).first_or_404()
    return jsonify({
        "id": draft.id, "documentType": draft.document_type, "title": draft.title,
        "content": draft.content, "clientName": draft.client_name,
        "oppositeName": draft.opposite_name, "createdAt": draft.created_at.isoformat(),
        "userId": draft.user_id,
    })

@app.route("/api/drafts/<int:draft_id>", methods=["DELETE"])
@require_auth
def delete_draft(draft_id):
    draft = GeneratedDraft.query.filter_by(id=draft_id, user_id=g.user.id).first_or_404()
    db.session.delete(draft)
    db.session.commit()
    return jsonify({"success": True})

@app.route("/api/drafts/<int:draft_id>/pdf", methods=["GET"])
@require_auth
def download_pdf(draft_id):
    draft = GeneratedDraft.query.filter_by(id=draft_id, user_id=g.user.id).first_or_404()
    pdf_bytes = build_pdf(draft.title, draft.content)
    filename = draft.title.replace(" ", "_").replace("/", "-")[:80] + ".pdf"
    return Response(
        pdf_bytes,
        mimetype="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

# ═══════════════════════════════════════════════════════════════════════════════
# Routes — Email (Gmail SMTP)
# ═══════════════════════════════════════════════════════════════════════════════

def log_email_attempt(user_id: int, status: str, recipient: str, subject: str, detail: str = ""):
    audit = {
        "status": status,
        "recipient": (recipient or "")[:254],
        "subject": (subject or "")[:500],
    }
    if detail:
        audit["detail"] = detail[:300]
    log_action(
        "email_sent" if status == "sent" else "email_failed",
        user_id,
        json.dumps(audit, ensure_ascii=False),
    )

def build_email_bodies(sender_name: str, sender_email: str, message: str, has_attachment: bool):
    plain_parts = [
        "AI LEGAL",
        "",
        "Sender Details",
        "",
        f"Name: {sender_name}",
        f"Email: {sender_email}",
        "",
        "Message",
        "",
        message,
    ]
    if has_attachment:
        plain_parts.extend(["", "Attachment", "", "The requested legal document is attached to this email."])
    text_body = "\n".join(plain_parts)

    safe_name = html.escape(sender_name, quote=True)
    safe_email = html.escape(sender_email, quote=True)
    safe_message = html.escape(message, quote=True).replace("\n", "<br>\n")
    attachment_html = (
        '<section style="margin-top:24px;padding:16px;background:#f5f7fa;border-radius:8px">'
        '<strong>Attachment</strong><p style="margin:8px 0 0">The requested legal document is attached to this email.</p></section>'
        if has_attachment else ""
    )
    html_body = (
        '<!doctype html><html><body style="margin:0;background:#f4f6f8;font-family:Arial,sans-serif;color:#172033">'
        '<main style="max-width:640px;margin:32px auto;background:#fff;border:1px solid #e4e8ef;border-radius:12px;overflow:hidden">'
        '<header style="padding:24px 32px;background:#101828;color:#fff;font-size:20px;font-weight:700;letter-spacing:.08em">AI LEGAL</header>'
        '<div style="padding:28px 32px">'
        '<h2 style="margin:0 0 16px;font-size:18px">Sender Details</h2>'
        f'<p style="margin:6px 0"><strong>Name:</strong> {safe_name}</p>'
        f'<p style="margin:6px 0"><strong>Email:</strong> {safe_email}</p>'
        '<hr style="border:0;border-top:1px solid #e4e8ef;margin:24px 0">'
        '<h2 style="margin:0 0 12px;font-size:18px">Message</h2>'
        f'<div style="line-height:1.65;white-space:normal">{safe_message}</div>'
        f'{attachment_html}'
        '</div><footer style="padding:16px 32px;background:#f8fafc;color:#667085;font-size:12px">'
        'Sent securely from the AI Legal website.</footer></main></body></html>'
    )
    return text_body, html_body

@app.route("/api/send-email", methods=["POST"])
@require_auth
def send_email():
    if request.content_length and request.content_length > MAX_EMAIL_REQUEST_BYTES:
        log_email_attempt(g.user.id, "failed", "", "", "Request exceeds 8 MB limit")
        return jsonify({"error": "The email attachment is too large. Maximum request size is 8 MB."}), 413

    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        log_email_attempt(g.user.id, "failed", "", "", "Invalid JSON request")
        return jsonify({"error": "Please enter recipient, sender details, subject, and message."}), 400

    def input_value(primary, *legacy):
        for key in (primary, *legacy):
            if key in data and data[key] is not None:
                return data[key]
        return ""

    recipient = str(input_value("to", "recipientEmail")).strip()
    sender_name = str(input_value("sender_name", "senderName")).strip()
    sender_email = str(input_value("sender_email", "senderEmail")).strip()
    subject = str(input_value("subject")).strip()
    body = str(input_value("message", "body")).strip()

    if not is_valid_email(recipient):
        log_email_attempt(g.user.id, "failed", recipient, subject, "Invalid recipient address")
        return jsonify({"error": "Please enter a valid recipient email address."}), 400
    if not sender_name or len(sender_name) > 200:
        log_email_attempt(g.user.id, "failed", recipient, subject, "Invalid sender name")
        return jsonify({"error": "Please enter a valid sender name."}), 400
    if not is_valid_email(sender_email):
        log_email_attempt(g.user.id, "failed", recipient, subject, "Invalid sender email address")
        return jsonify({"error": "Please enter a valid sender email address."}), 400
    if not subject:
        log_email_attempt(g.user.id, "failed", recipient, subject, "Missing subject")
        return jsonify({"error": "Please enter an email subject."}), 400
    if not body:
        log_email_attempt(g.user.id, "failed", recipient, subject, "Missing message")
        return jsonify({"error": "Please enter a message."}), 400

    draft_id = input_value("document_id", "documentId", "draftId") or None
    if isinstance(draft_id, str) and draft_id.isdigit():
        draft_id = int(draft_id)

    attachments = []
    attachment = data.get("attachment")
    if attachment is not None:
        try:
            attachments.append(validate_email_attachment(attachment))
        except (OverflowError, ValueError) as exc:
            log_email_attempt(g.user.id, "failed", recipient, subject, str(exc))
            return jsonify({"error": str(exc)}), 400

    if draft_id is not None and isinstance(draft_id, int):
        draft = GeneratedDraft.query.filter_by(id=draft_id, user_id=g.user.id).first()
        if not draft:
            log_email_attempt(g.user.id, "failed", recipient, subject, "Generated document not found")
            return jsonify({"error": "The selected document could not be found."}), 404
        try:
            pdf_bytes = build_pdf(draft.title, draft.content)
            filename = secure_filename(draft.title) or "legal-document"
            attachments.insert(0, {
                "filename": f"{filename[:180]}.pdf",
                "content": base64.b64encode(pdf_bytes).decode("ascii"),
            })
        except Exception:
            return jsonify({"error": "The generated document could not be prepared as a PDF."}), 500

    text_body, html_body = build_email_bodies(sender_name, sender_email, body, bool(attachments))
    email_rec = SentEmail(
        recipient_email=recipient,
        sender_name=sender_name,
        sender_email=sender_email,
        subject=subject,
        body=body,
        draft_id=draft_id,
        status="pending",
        user_id=g.user.id,
    )
    db.session.add(email_rec)
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        return jsonify({"error": "The email record could not be saved."}), 500

    try:
        email_rec.message_id = send_email_smtp(recipient, subject, text_body, attachments, html_body)
    except EmailConfigurationError as exc:
        email_rec.status = "failed"
        db.session.commit()
        log_email_attempt(g.user.id, "failed", recipient, subject, str(exc))
        return jsonify({"error": str(exc)}), 503
    except EmailProviderError as exc:
        email_rec.status = "failed"
        db.session.commit()
        log_email_attempt(g.user.id, "failed", recipient, subject, str(exc))
        return jsonify({"error": str(exc)}), exc.status_code

    email_rec.status = "sent"
    email_rec.sent_at = datetime.datetime.utcnow()
    db.session.commit()
    log_email_attempt(g.user.id, "sent", recipient, subject)

    return jsonify({
        "id": email_rec.id,
        "recipientEmail": email_rec.recipient_email,
        "senderName": email_rec.sender_name,
        "senderEmail": email_rec.sender_email,
        "subject": email_rec.subject,
        "body": email_rec.body,
        "draftId": email_rec.draft_id,
        "documentId": email_rec.draft_id,
        "sentAt": email_rec.sent_at.isoformat(),
        "status": email_rec.status,
        "messageId": email_rec.message_id,
    })

@app.route("/api/emails", methods=["GET"])
@require_auth
def list_emails():
    emails = SentEmail.query.filter_by(user_id=g.user.id).all()
    return jsonify([{
        "id": e.id,
        "recipientEmail": e.recipient_email,
        "senderName": e.sender_name,
        "senderEmail": e.sender_email,
        "subject": e.subject,
        "body": e.body,
        "draftId": e.draft_id,
        "documentId": e.draft_id,
        "sentAt": e.sent_at.isoformat(),
        "status": e.status,
        "messageId": e.message_id,
    } for e in emails])

# ═══════════════════════════════════════════════════════════════════════════════
# Routes — History & Dashboard
# ═══════════════════════════════════════════════════════════════════════════════

@app.route("/api/history", methods=["GET"])
@require_auth
def history():
    uid = g.user.id
    docs   = Document.query.filter_by(user_id=uid).limit(5).all()
    chats  = ChatHistory.query.filter_by(user_id=uid, role="user").limit(5).all()
    drafts = GeneratedDraft.query.filter_by(user_id=uid).limit(5).all()
    emails = SentEmail.query.filter_by(user_id=uid).limit(5).all()

    items = (
        [{"id": d.id,        "type": "upload",   "description": f"Uploaded: {d.name}",                     "createdAt": d.upload_date.isoformat(),  "metadata": None} for d in docs]   +
        [{"id": c.id * 100,  "type": "chat",     "description": f"Asked: {c.content[:60]}...",              "createdAt": c.created_at.isoformat(),   "metadata": None} for c in chats]  +
        [{"id": d.id * 300,  "type": "generate", "description": f"Generated: {d.title}",                   "createdAt": d.created_at.isoformat(),   "metadata": None} for d in drafts] +
        [{"id": e.id * 400,  "type": "email",    "description": f"Sent email to: {e.recipient_email}",     "createdAt": e.sent_at.isoformat(),      "metadata": None} for e in emails]
    )
    items.sort(key=lambda x: x["createdAt"], reverse=True)
    return jsonify(items)

@app.route("/api/dashboard/summary", methods=["GET"])
@require_auth
def dashboard_summary():
    uid = g.user.id
    doc_count   = Document.query.filter_by(user_id=uid).count()
    chat_count  = ChatHistory.query.filter_by(user_id=uid).count()
    draft_count = GeneratedDraft.query.filter_by(user_id=uid).count()
    email_count = SentEmail.query.filter_by(user_id=uid, status="sent").count()

    recent_docs   = Document.query.filter_by(user_id=uid).limit(3).all()
    recent_chats  = ChatHistory.query.filter_by(user_id=uid, role="user").limit(3).all()
    recent_drafts = GeneratedDraft.query.filter_by(user_id=uid).limit(2).all()

    activity = sorted([
        *[{"id": d.id,       "type": "upload",   "description": f"Uploaded: {d.name}",        "createdAt": d.upload_date.isoformat(), "metadata": None} for d in recent_docs],
        *[{"id": c.id * 100, "type": "chat",     "description": f"Asked: {c.content[:50]}...", "createdAt": c.created_at.isoformat(),  "metadata": None} for c in recent_chats],
        *[{"id": d.id * 300, "type": "generate", "description": f"Generated: {d.title}",       "createdAt": d.created_at.isoformat(),  "metadata": None} for d in recent_drafts],
    ], key=lambda x: x["createdAt"], reverse=True)[:5]

    return jsonify({
        "uploadedDocuments": doc_count,
        "questionsAsked":    chat_count,
        "draftsGenerated":   draft_count,
        "advocatesFound":    12,
        "lettersSent":       email_count,
        "recentActivity":    activity,
    })

# ═══════════════════════════════════════════════════════════════════════════════
# Routes — Admin
# ═══════════════════════════════════════════════════════════════════════════════

@app.route("/api/admin/stats", methods=["GET"])
@require_auth
def admin_stats():
    logs = AdminLog.query.order_by(AdminLog.created_at.desc()).limit(20).all()
    return jsonify({
        "totalUsers":     User.query.count(),
        "totalDocuments": Document.query.count(),
        "totalQueries":   ChatHistory.query.count(),
        "generatedPdfs":  GeneratedDraft.query.count(),
        "sentEmails":     SentEmail.query.filter_by(status="sent").count(),
        "totalDatasets":  Dataset.query.count(),
        "systemActivity": [{
            "id": l.id, "action": l.action, "userId": l.user_id,
            "details": l.details, "createdAt": l.created_at.isoformat(),
        } for l in logs],
    })

@app.route("/api/admin/users", methods=["GET"])
@require_auth
def admin_users():
    users = User.query.all()
    return jsonify([{
        "id": u.id, "name": u.name, "email": u.email,
        "role": u.role, "createdAt": u.created_at.isoformat(),
    } for u in users])

@app.route("/api/admin/logs", methods=["GET"])
@require_auth
def admin_logs():
    logs = AdminLog.query.order_by(AdminLog.created_at.desc()).all()
    return jsonify([{
        "id": l.id, "action": l.action, "userId": l.user_id,
        "details": l.details, "createdAt": l.created_at.isoformat(),
    } for l in logs])

@app.route("/api/admin/reindex", methods=["POST"])
@require_auth
def admin_reindex():
    """Re-index all documents into ChromaDB."""
    docs = Document.query.filter(Document.content.isnot(None)).all()
    count = 0
    for doc in docs:
        if doc.content:
            index_document(doc.id, doc.content)
            count += 1
    log_action("reindex_documents", g.user.id, f"Re-indexed {count} documents into ChromaDB")
    return jsonify({"success": True, "message": f"Re-indexed {count} documents into ChromaDB vector store."})

# ═══════════════════════════════════════════════════════════════════════════════
# Routes — Advocates (Search)
# ═══════════════════════════════════════════════════════════════════════════════

CITY_PRESETS = {
    "hanamkonda": {
        "lat": 18.0069, "lon": 79.5579, "label": "Hanamkonda, Hanumakonda, Telangana, India",
        "areas": ["Court Circle", "Subedari", "Balasamudram", "Naimnagar", "Kazipet Main Road", "Chowrashta", "Hanamkonda Bus Stand Area", "Lashkar Bazaar", "Kakatiya University Road"]
    },
    "warangal": {
        "lat": 17.9689, "lon": 79.5941, "label": "Warangal, Telangana, India",
        "areas": ["Station Road", "MGM Hospital Circle", "Girmajipet", "Kashibugga", "District Court Complex", "Enumamula", "Under Bridge Road"]
    },
    "mumbai": {
        "lat": 19.0760, "lon": 72.8777, "label": "Mumbai, Maharashtra, India",
        "areas": ["Fort", "Nariman Point", "Bandra Kurla Complex", "Marine Lines Court Area", "Dadar West", "Ballard Estate", "Andheri East", "Colaba", "Vile Parle"]
    },
    "hyderabad": {
        "lat": 17.3850, "lon": 78.4867, "label": "Hyderabad, Telangana, India",
        "areas": ["Banjara Hills", "Nampally District Court Area", "Jubilee Hills", "High Court Road", "Hitec City", "Abids", "Secunderabad", "Madhapur", "Himayatnagar"]
    },
    "delhi": {
        "lat": 28.6139, "lon": 77.2090, "label": "New Delhi, Delhi, India",
        "areas": ["Tilak Marg (Supreme Court Area)", "Patiala House Court Complex", "Connaught Place", "Saket District Court Area", "Rohini Court Complex", "Lajpat Nagar", "Barakhamba Road", "Dwarka Sector 10 Court Area"]
    },
    "bengaluru": {
        "lat": 12.9716, "lon": 77.5946, "label": "Bengaluru, Karnataka, India",
        "areas": ["MG Road", "Indiranagar", "Koramangala", "High Court / Cubbon Park Area", "Jayanagar", "Malleshwaram", "Whitefield", "Rajajinagar"]
    },
    "chennai": {
        "lat": 13.0827, "lon": 80.2707, "label": "Chennai, Tamil Nadu, India",
        "areas": ["High Court Buildings (George Town)", "Parrys Corner", "T. Nagar", "Anna Nagar", "Mylapore", "Egmore Court Area", "Nungambakkam"]
    },
    "kolkata": {
        "lat": 22.5726, "lon": 88.3639, "label": "Kolkata, West Bengal, India",
        "areas": ["High Court Area (Council House St)", "BBD Bagh", "Park Street", "Salt Lake Sector V", "Alipore Court Area", "Howrah Court Circle"]
    },
    "pune": {
        "lat": 18.5204, "lon": 73.8567, "label": "Pune, Maharashtra, India",
        "areas": ["Shivajinagar District Court Circle", "FC Road", "Kothrud", "Viman Nagar", "Camp Area", "Deccan Gymkhana", "Hadapsar"]
    },
    "ahmedabad": {
        "lat": 23.0225, "lon": 72.5714, "label": "Ahmedabad, Gujarat, India",
        "areas": ["High Court Complex (Sola)", "Ashram Road", "CG Road", "Bodakdev", "Navrangpura", "Mirzapur Court Area"]
    }
}

LAWYER_TEMPLATES = [
    {"name": "{prefix} {surname} & Associates", "type": "Law Firm", "specs": ["Civil Law", "Corporate / Business Law"], "phone_prefix": "+91 98"},
    {"name": "Advocate {given_name} {surname}", "type": "Lawyer", "specs": ["Criminal Law"], "phone_prefix": "+91 94"},
    {"name": "{surname} Legal Chambers", "type": "Law Firm", "specs": ["Property / Real Estate Law", "Civil Law"], "phone_prefix": "+91 97"},
    {"name": "Advocate {given_name} {surname} (High Court)", "type": "High Court Advocate", "specs": ["Constitutional Law", "Civil Law"], "phone_prefix": "+91 99"},
    {"name": "{surname} & {surname} Family Law Associates", "type": "Law Firm", "specs": ["Family Law", "Divorce / Matrimonial"], "phone_prefix": "+91 93"},
    {"name": "Advocate {given_name} {surname}", "type": "Lawyer", "specs": ["Cyber / IT Law", "Intellectual Property"], "phone_prefix": "+91 91"},
    {"name": "{surname} Tax & Commercial Consultancy", "type": "Legal Consultancy", "specs": ["Tax Law", "Corporate / Business Law"], "phone_prefix": "+91 98"},
    {"name": "Advocate {given_name} {surname}", "type": "Lawyer", "specs": ["Consumer Law", "Civil Law"], "phone_prefix": "+91 96"},
    {"name": "Legal Aid & Advocate {given_name} {surname}", "type": "Lawyer", "specs": ["Employment / Labour Law", "Human Rights Law"], "phone_prefix": "+91 95"},
    {"name": "{given_name} {surname} Notary & Advocate", "type": "Notary / Advocate", "specs": ["Property / Real Estate Law", "Civil Law"], "phone_prefix": "+91 92"},
    {"name": "{surname} Corporate Law Practice", "type": "Law Firm", "specs": ["Corporate / Business Law", "Banking & Finance Law"], "phone_prefix": "+91 90"},
    {"name": "Advocate {given_name} {surname} (District Court)", "type": "District Court Advocate", "specs": ["Criminal Law", "Civil Law"], "phone_prefix": "+91 94"},
    {"name": "{prefix} {given_name} {surname} Chambers", "type": "Law Office", "specs": ["Immigration Law", "Civil Law"], "phone_prefix": "+91 97"},
    {"name": "Advocate {given_name} {surname} & Partners", "type": "Law Firm", "specs": ["Environmental Law", "Property / Real Estate Law"], "phone_prefix": "+91 98"}
]

GIVEN_NAMES = ["Rajesh", "Srinivas", "Ananya", "Ramesh", "Priya", "Venkatesh", "Kavitha", "Suresh", "Lakshmi", "Amit", "Vikram", "Sneha", "Prashanth", "Deepak", "Sunita", "Ravi", "Gautam", "Aparna", "Arjun", "Neha"]
SURNAMES = ["Rao", "Reddy", "Sharma", "Verma", "Joshi", "Kulkarni", "Patel", "Gupta", "Nair", "Iyer", "Chowdary", "Deshmukh", "Singhania", "Mehta", "Kumar", "Singh", "Banerjee", "Chatterjee"]
PREFIXES = ["V.", "K.", "M.", "S.", "P.", "R.", "A.", "G.", "T.", "N."]

def geocode_location(city_name=None, lat=None, lon=None):
    if lat is not None and lon is not None:
        try:
            url = f"https://nominatim.openstreetmap.org/reverse?format=json&lat={lat}&lon={lon}"
            req = Request(url, headers={'User-Agent': 'LegalGPT-App/1.0'})
            with urlopen(req, timeout=4) as resp:
                data = json.loads(resp.read().decode())
                label = data.get("display_name") or f"Location ({lat:.4f}, {lon:.4f})"
                return float(lat), float(lon), label
        except Exception:
            return float(lat), float(lon), f"Location ({float(lat):.4f}, {float(lon):.4f})"

    if not city_name:
        city_name = "Hanamkonda"

    city_clean = city_name.strip().lower()
    for key, preset in CITY_PRESETS.items():
        if key in city_clean:
            return preset["lat"], preset["lon"], preset["label"]

    try:
        url = f"https://nominatim.openstreetmap.org/search?format=json&q={urlencode({'q': city_name})[2:]}&limit=1"
        req = Request(url, headers={'User-Agent': 'LegalGPT-App/1.0'})
        with urlopen(req, timeout=5) as resp:
            data = json.loads(resp.read().decode())
            if data:
                return float(data[0]['lat']), float(data[0]['lon']), data[0].get('display_name', city_name.strip().title())
    except Exception as exc:
        app.logger.warning("Nominatim geocode failed for %s: %s", city_name, exc)

    return 18.0069, 79.5579, f"{city_name.strip().title()}, India"

def generate_city_advocates(lat, lon, location_label, radius_m):
    seed_val = int((abs(lat) * 1000 + abs(lon) * 10000))
    rng = random.Random(seed_val)

    city_lower = location_label.lower()
    matched_preset = None
    for key, preset in CITY_PRESETS.items():
        if key in city_lower:
            matched_preset = preset
            break

    areas = matched_preset["areas"] if matched_preset else [
        "Court Road", "Civil Lines", "Main Market Area", "Station Circle", "District Court Complex", "High Court Road", "Collectorate Area", "Gandhi Circle", "Naya Bazaar", "Commercial Hub"
    ]

    city_clean_name = location_label.split(",")[0].strip()

    results = []
    count = rng.randint(12, 16)
    radius_deg = (radius_m / 111000.0) * 0.65

    for i in range(count):
        tmpl = LAWYER_TEMPLATES[i % len(LAWYER_TEMPLATES)]
        given = rng.choice(GIVEN_NAMES)
        surname = rng.choice(SURNAMES)
        prefix = rng.choice(PREFIXES)
        area = rng.choice(areas)

        name = tmpl["name"].format(prefix=prefix, given_name=given, surname=surname)
        spec = rng.choice(tmpl["specs"])
        phone = f"{tmpl['phone_prefix']}{rng.randint(10, 99)} {rng.randint(10000, 99999)}"
        domain_name = f"{given.lower()}{surname.lower()}legal.in".replace(" ", "")
        website = f"https://www.{domain_name}"
        hours = rng.choice(["Mo-Sa 09:30-18:30", "Mo-Fr 10:00-19:00", "Mo-Sa 10:00-20:00", "Mo-Fr 09:00-17:30"])

        angle = rng.uniform(0, 2 * math.pi)
        dist = rng.uniform(0.08, 0.95) * radius_deg
        adv_lat = round(lat + dist * math.cos(angle), 5)
        adv_lon = round(lon + dist * math.sin(angle), 5)

        address = f"Opp. {area}, {city_clean_name}"

        results.append({
            "id": f"adv_{i+1}_{seed_val}",
            "name": name,
            "address": address,
            "lat": adv_lat,
            "lon": adv_lon,
            "osmType": tmpl["type"],
            "phone": phone,
            "website": website,
            "openingHours": hours,
            "detectedSpecialization": spec
        })

    return results

@app.route("/api/advocates", methods=["GET"])
@require_auth
def advocates():
    return jsonify([])

@app.route("/api/advocates/search", methods=["POST"])
@require_auth
def search_nearby_advocates():
    data = request.get_json(silent=True) or {}
    city = data.get("city")
    lat_in = data.get("lat")
    lon_in = data.get("lon")
    radius_m = data.get("radiusMeters") or 10000

    try:
        radius_m = int(radius_m)
    except (TypeError, ValueError):
        radius_m = 10000

    center_lat, center_lon, location_label = geocode_location(
        city_name=city,
        lat=float(lat_in) if lat_in is not None else None,
        lon=float(lon_in) if lon_in is not None else None
    )

    results = generate_city_advocates(center_lat, center_lon, location_label, radius_m)

    return jsonify({
        "centerLat": center_lat,
        "centerLon": center_lon,
        "locationLabel": location_label,
        "results": results
    })

# ═══════════════════════════════════════════════════════════════════════════════
# Healthcheck
# ═══════════════════════════════════════════════════════════════════════════════

@app.route("/api/healthz", methods=["GET"])
def healthz():
    return jsonify({
        "status": "ok",
        "stack": "Flask + SQLAlchemy + ChromaDB + LangChain + NVIDIA NIM",
        "llm": f"{NVIDIA_MODEL} via LangChain NVIDIA NIM",
        "rag": "ChromaDB",
        "db": "SQLite",
        "emailConfigured": smtp_is_configured(),
    })

# ═══════════════════════════════════════════════════════════════════════════════
# Database Initialization & Seeding
# ═══════════════════════════════════════════════════════════════════════════════

def seed_admin():
    if not User.query.filter_by(email="admin@legalgpt.com").first():
        admin = User(
            name="Admin",
            email="admin@legalgpt.com",
            password_hash=hash_password("admin123"),
            role="admin",
        )
        db.session.add(admin)
        db.session.commit()
        app.logger.info("Admin user seeded: admin@legalgpt.com / admin123")

def seed_demo_user():
    user1 = User.query.filter_by(email="user@legalgpt.com").first()
    if not user1:
        user1 = User(
            name="Legalgpt",
            email="user@legalgpt.com",
            password_hash=hash_password("legalgpt@123"),
            role="user",
        )
        db.session.add(user1)
    else:
        user1.name = "Legalgpt"
        user1.password_hash = hash_password("legalgpt@123")

    user2 = User.query.filter_by(email="legalgpt@legalgpt.com").first()
    if not user2:
        user2 = User(
            name="Legalgpt",
            email="legalgpt@legalgpt.com",
            password_hash=hash_password("legalgpt@123"),
            role="user",
        )
        db.session.add(user2)
    else:
        user2.name = "Legalgpt"
        user2.password_hash = hash_password("legalgpt@123")

    db.session.commit()
    app.logger.info("Legalgpt user seeded/updated with password legalgpt@123")

with app.app_context():
    db.create_all()
    seed_admin()
    seed_demo_user()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8080))
    app.run(host="0.0.0.0", port=port, debug=False)
